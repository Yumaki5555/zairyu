# 在留資格「定住者」の在留目的（日系３世・配偶者・未成年未婚の実子・その他）ごとの人数を、国の表から1行に1つの数の表にして確かめる
# 使い方: <pandas・openpyxl の入った python> -I make_teiju.py
#
#   読むもの（どれも download_estat.py・download_nyukoku.py が落とした e-Stat の表と、それから作った データ/ の表）
#     2023年12月末から  在留外国人統計「国籍・地域別 在留目的別 在留外国人（特定活動・日本人の配偶者等・定住者）」（〇-02m.xlsx）。
#                     定住者の在留目的が分かれているのはこの表から（2023年6月末までの表は「定住者」の合計だけ）
#     新しく入国した人  出入国管理統計「国籍・地域別 入国目的別 新規入国外国人（短期滞在・特定活動等）」（元データ/出入国管理統計/。1年間の数）の定住者の列
#   書き出す表（データ/定住者/）
#     在留目的別.csv    列：時点・在留目的・人数・出どころ（全国。在留目的は国の表の名前から（告示〇号）を除いたもの。「総数」は定住者の合計）
#     国籍_在留目的.csv 列：時点・国籍・地域・在留目的・人数・出どころ（0人の行は書かない。「うち中国〔香港〕」などの内わけの行は入れない）
#     新規入国_在留目的.csv 列：年・国籍・地域・在留目的・告示の号・人数・出どころ（「総数」の行は全国。州の行は入れない。0人の行は書かない）
#   確かめ: データ/定住者/確かめ.csv
#     時点ごとに、在留目的の合計＝総数、総数＝データ/全国_国籍_在留資格.csv の定住者の数、国籍ごとの合計＝全国の数。
#     新しく入国した人は、在留目的の合計＝総数、国籍ごとの合計＝全国の数
import csv
import os
import re
import sys
from collections import defaultdict

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "データ")
SRC = os.path.join(BASE, "元データ")
OUT = os.path.join(DATA, "定住者")
TOTAL = "総数"
VISA = "定住者"
REGIONS = {"アジア", "ヨーロッパ", "アフリカ", "北アメリカ", "南アメリカ", "オセアニア"}  # 新しく入国した人の表の州の行
ENTRY_LIST = os.path.join(SRC, "出入国管理統計_出どころ一覧.csv")
Z2H = str.maketrans("０１２３４５６７８９", "0123456789")
CHECKS = []


def norm(s):
    return re.sub(r"[\s　]+", "", str(s)) if s is not None else ""


def to_int(v):
    s = norm(v).replace(",", "")
    return int(float(s)) if re.fullmatch(r"\d+(\.0+)?", s) else 0


def purpose_of(s):
    """「日系３世（告示３・４号）」→（日系３世, 3号・4号）。2023年12月末の表と新しく入国した人の表は見出しに告示の号がある"""
    s = norm(s)
    m = re.search(r"（告示(.+?)）", s)
    no = ""
    if m:
        nums = re.findall(r"\d+", m.group(1).translate(Z2H))
        no = "・".join(f"{n}号" for n in nums)
    name = re.sub(r"（告示.+?）", "", s)
    return (TOTAL if name in ("計", TOTAL) else name), no


def read_long(path):
    """在留外国人統計の第2表（1行に1つの数）：全国の在留目的ごとの数と、国籍ごとの数"""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    japan, by_nat = {}, defaultdict(dict)
    for r in ws.iter_rows(values_only=True):
        if len(r) < 6 or norm(r[3]) != VISA:
            continue
        area, nat, (purpose, _), v = norm(r[1]), norm(r[2]), purpose_of(r[4]), to_int(r[5])
        if area == TOTAL and nat == TOTAL:
            japan[purpose] = v
        elif area == "無国籍" and nat == TOTAL:  # 無国籍は州の行だけ
            by_nat["無国籍"][purpose] = v
        elif area != TOTAL and nat != TOTAL and not nat.startswith("うち"):
            by_nat[nat][purpose] = v
    return japan, by_nat


def read_entry(path):
    """新しく入国した人の表：見出しから定住者の列（総数・在留目的）を探し、国籍・地域ごとの数を返す → {国籍: {(在留目的, 告示の号): 人数}}"""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    first = next(i for i, r in enumerate(rows) if norm(r[0]) == TOTAL)
    head = [[norm(c) for c in r] for r in rows[:first]]
    top = next(i for i, r in enumerate(head) if VISA in r)
    c0 = head[top].index(VISA)
    c1 = next((c for c in range(c0 + 1, len(head[top])) if head[top][c]), len(head[top]))
    cols = {c: purpose_of(head[top + 1][c]) for c in range(c0, c1)}
    out = {}
    for r in rows[first:]:
        n = norm(r[0])
        if not n or n in REGIONS or n.startswith("うち"):
            continue
        out[n] = {k: to_int(r[c]) for c, k in cols.items()}
    return out


def main():
    os.makedirs(OUT, exist_ok=True)
    rows, nat_rows = [], []
    nat = pd.read_csv(os.path.join(DATA, "全国_国籍_在留資格.csv"))
    ref = {w: int(v) for w, v in nat[nat["在留資格"] == VISA].groupby("時点")["人数"].sum().items()}  # 時点 → 定住者の数（ほかの表）
    for r in csv.DictReader(open(os.path.join(SRC, "出どころ一覧.csv"), encoding="utf-8-sig")):
        if "在留目的別在留外国人（特定活動" not in r["表題"]:
            continue
        w, f = r["時点"], r["ファイル"]
        japan, by_nat = read_long(os.path.join(SRC, f))
        rows += [[w, p, v, f] for p, v in japan.items()]
        nat_rows += [[w, n, p, v, f] for n, d in sorted(by_nat.items()) for p, v in d.items() if v and p != TOTAL]
        s = sum(v for p, v in japan.items() if p != TOTAL)
        CHECKS.append([w, "在留目的の合計＝総数", s, japan[TOTAL], s - japan[TOTAL]])
        CHECKS.append([w, "総数＝在留資格の表の定住者", japan[TOTAL], ref[w], japan[TOTAL] - ref[w]])
        for p in japan:  # 国籍ごとの合計＝全国の数
            s = sum(d.get(p, 0) for d in by_nat.values())
            CHECKS.append([w, f"国籍の合計（{p}）", s, japan[p], s - japan[p]])
    # 新しく入国した人（1年間）
    entry_rows = []
    for r in csv.DictReader(open(ENTRY_LIST, encoding="utf-8-sig")):
        f = "出入国管理統計/" + r["ファイル"]
        got = read_entry(os.path.join(SRC, f))
        japan = got.pop(TOTAL)
        entry_rows += [[r["年"], TOTAL, p, no, v, f] for (p, no), v in japan.items()]
        entry_rows += [[r["年"], n, p, no, v, f] for n, d in got.items() for (p, no), v in d.items() if v]
        total = next(v for (p, _), v in japan.items() if p == TOTAL)
        s = sum(v for (p, _), v in japan.items() if p != TOTAL)
        CHECKS.append([r["年"], "新しく入国した人：在留目的の合計＝総数", s, total, s - total])
        for k, v in japan.items():
            s = sum(d.get(k, 0) for d in got.values())
            CHECKS.append([r["年"], f"新しく入国した人：国籍の合計（{k[0]}）", s, v, s - v])
    pd.DataFrame(rows, columns=["時点", "在留目的", "人数", "出どころ"]).to_csv(os.path.join(OUT, "在留目的別.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(nat_rows, columns=["時点", "国籍・地域", "在留目的", "人数", "出どころ"]).to_csv(os.path.join(OUT, "国籍_在留目的.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(entry_rows, columns=["年", "国籍・地域", "在留目的", "告示の号", "人数", "出どころ"]).to_csv(os.path.join(OUT, "新規入国_在留目的.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(CHECKS, columns=["時点", "確かめ", "数", "比べる数", "差"]).to_csv(os.path.join(OUT, "確かめ.csv"), index=False, encoding="utf-8-sig")
    bad = [c for c in CHECKS if c[4]]
    print(f"在留目的別 {len(rows)} 行、国籍_在留目的 {len(nat_rows)} 行、新規入国_在留目的 {len(entry_rows)} 行")
    print("確かめ：" + ("すべて合っています" if not bad else f"合わない所 {len(bad)}"))
    for c in bad:
        print("  ", c)


if __name__ == "__main__":
    main()
