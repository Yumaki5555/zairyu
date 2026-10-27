# 在留資格「特定活動」の在留目的（活動の中身）ごとの人数を、国の表から1行に1つの数の表にして確かめる
# 使い方: <pandas・openpyxl の入った python> -I make_tokutei.py
#
#   読むもの（どれも download_estat.py が落とした e-Stat の表と、それから作った データ/ の表）
#     1990〜2005年末  データ/古い表/全国_国籍_在留資格.csv の「総数」の行（統計書の「特定活動」と内わけ。内わけ1＝ワーキング・ホリデー、内わけ2＝その他）
#     2006〜2011年末  データ/全国_国籍_在留資格.csv（国籍ごとの数を足す。内わけはワーキング・ホリデーとその他）
#     2012年12月末〜2022年6月末  「国籍・地域別 在留資格（在留目的）別 在留外国人」（元データ/在留外国人統計/〇/〇-01-1.xlsx）の「総数」の行。
#                     「特定活動」の列の下に在留目的の列があり、いくつかはさらに「本人」「家族」に分かれる（データ/全国_国籍_在留資格.csv は
#                     「家族」の列を読んでいないので、ここでは表から読み直す。特定活動の合計の数は同じ）
#     2022年12月末・2023年6月末  同じ表（〇-01m.xlsx。1行に1つの数の形）
#     2023年12月末から  「国籍・地域別 在留目的別 在留外国人（特定活動・日本人の配偶者等・定住者）」（〇-02m.xlsx）
#     新しく入国した人  出入国管理統計「国籍・地域別 入国目的別 新規入国外国人（短期滞在・特定活動等）」（download_nyukoku.py が
#                     元データ/出入国管理統計/ に落とした表。1年間の数）の「総数」の行の、特定活動の列
#   書き出す表（データ/特定活動/）
#     在留目的別.csv    列：時点・統計・在留目的・人数・出どころ（全国。在留目的は国の表の名前のまま、空白は除く。「総数」は特定活動の合計）
#     国籍_在留目的.csv 列：時点・国籍・地域・在留目的・人数・出どころ（2022年12月末から。0人の行は書かない。「うち中国〔香港〕」などの内わけの行は入れない）
#     新規入国_入国目的.csv 列：年・入国目的・告示の号・人数・出どころ（全国。入国目的は見出しの2段を「アマチュアスポーツ選手（本人）」のようにつなぐ。
#                     告示の号は見出しの（告示〇号）を「1号・2号の2」のように書く。「総数」は特定活動の合計で、デジタルノマドも入る）
#   確かめ: データ/特定活動/確かめ.csv
#     時点ごとに、在留目的の合計＝総数（デジタルノマドは中長期在留者ではないので総数の外。表の注のとおり）、
#     総数＝データ/全国_国籍_在留資格.csv（または古い表）の特定活動の数、国籍ごとの合計＝全国の数。新しく入国した人は、入国目的の合計＝総数
import csv
import os
import re
import sys
from collections import defaultdict

import openpyxl
import pandas as pd

sys.stdout.reconfigure(encoding="utf-8")
BASE = os.path.dirname(os.path.abspath(__file__))
DATA = os.path.join(BASE, "データ")
SRC = os.path.join(BASE, "元データ")
OUT = os.path.join(DATA, "特定活動")
TOTAL = "総数"
OUTSIDE = ("デジタルノマド",)  # 総数の外（外数）の在留目的
NEXT_VISA = {"永住者", "日本人の配偶者等", "永住者の配偶者等", "定住者", "特別永住者"}  # 表で特定活動の右にある在留資格
ENTRY_LIST = os.path.join(SRC, "出入国管理統計_出どころ一覧.csv")
Z2H = str.maketrans("０１２３４５６７８９", "0123456789")
CHECKS = []


def norm(s):
    return re.sub(r"[\s　]+", "", str(s)) if s is not None else ""


def to_int(v):
    s = norm(v).replace(",", "")
    return int(float(s)) if re.fullmatch(r"\d+(\.0+)?", s) else 0


def sources():
    """時点 → (表の種類, ファイル)。出どころ一覧.csv から国籍・地域別の在留目的の表を選ぶ"""
    out = {}
    for r in csv.DictReader(open(os.path.join(SRC, "出どころ一覧.csv"), encoding="utf-8-sig")):
        w, t, f = r["時点"], r["表題"], r["ファイル"]
        if w < "2012" or "総在留" in t:
            continue
        if "在留目的別在留外国人（特定活動" in t:
            out[w] = ("02m", f)
        elif "在留資格（在留目的）別在留外国人" in t and w not in out:
            out[w] = ("01m" if f.endswith("m.xlsx") else "01-1", f)
    return out


def read_wide(path):
    """2012〜2022年6月末の表：見出しの行から「特定活動」の列の範囲と在留目的の名前を作り、「総数」の行の数を読む"""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    first = next(i for i, r in enumerate(rows) if norm(r[0]) == TOTAL)
    head = [[norm(c) for c in r] for r in rows[:first]]
    top = next(i for i, r in enumerate(head) if "特定活動" in r)
    c0 = head[top].index("特定活動")
    # 特定活動の列は、次の在留資格の見出しの前まで（2020年12月末などは「永住者」の見出しが1段下にある）
    c1 = next((c for c in range(c0 + 1, len(head[top])) if head[top][c] or head[top + 1][c] in NEXT_VISA), len(head[top]))
    names, l2 = {}, ""
    for c in range(c0, c1):
        l2 = head[top + 1][c] or l2 if top + 1 < first else ""
        l3 = head[top + 2][c] if top + 2 < first else ""
        names[c] = TOTAL if l2 == "計" else l2 + l3
    return {n: to_int(rows[first][c]) for c, n in names.items()}


def read_long(path, kind):
    """2022年12月末からの表（1行に1つの数）：全国の在留目的ごとの数と、国籍ごとの数"""
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
    japan, by_nat = {}, defaultdict(dict)
    for r in ws.iter_rows(values_only=True):
        if len(r) < 6 or norm(r[3]) != "特定活動":
            continue
        area, nat, purpose, v = norm(r[1]), norm(r[2]), norm(r[4]), to_int(r[5])
        purpose = TOTAL if purpose in ("計", TOTAL) else purpose
        if area == TOTAL and nat == TOTAL:
            japan[purpose] = v
        elif area == "無国籍" and nat == TOTAL:  # 無国籍は州の行だけ
            by_nat["無国籍"][purpose] = v
        elif area != TOTAL and nat != TOTAL and not nat.startswith("うち"):
            by_nat[nat][purpose] = v
    return japan, by_nat


def read_entry(path):
    """新しく入国した人の表：見出しの2段（入国目的・本人などの分け方）と告示の号を読み、「総数」の行の特定活動の列の数を返す"""
    rows = list(openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True))
    first = next(i for i, r in enumerate(rows) if norm(r[0]) == TOTAL)
    head = [[norm(c) for c in r] for r in rows[:first]]
    top = next(i for i, r in enumerate(head) if "特定活動" in r)
    c0 = head[top].index("特定活動")
    c1 = next((c for c in range(c0 + 1, len(head[top])) if head[top][c]), len(head[top]))
    out, l2 = [], ""
    for c in range(c0, c1):
        l2 = head[top + 1][c] or l2
        l3 = head[top + 2][c] if top + 2 < first else ""
        no = re.findall(r"（告示(.+?)）", l3 or l2)
        name = re.sub(r"（告示.+?）", "", l2) + (f"（{re.sub(r'（告示.+?）', '', l3)}）" if l3 else "")
        out.append((name, no[0].replace("告示", "").translate(Z2H).replace("、", "・") if no else "", to_int(rows[first][c])))
    return out


def main():
    os.makedirs(OUT, exist_ok=True)
    rows, nat_rows = [], []
    # 1990〜2005年末（古い表）
    old = pd.read_csv(os.path.join(DATA, "古い表", "全国_国籍_在留資格.csv"))
    old = old[(old["国籍・地域"] == TOTAL) & old["在留資格"].str.startswith("特定活動")]
    sub = {"特定活動": TOTAL, "特定活動（内わけ1）": "ワーキング・ホリデー", "特定活動（内わけ2）": "その他"}
    ref = {}  # 時点 → 特定活動の数（確かめに使う、ほかの表の数）
    for w, v, n, f in zip(old["時点"], old["在留資格"], old["人数"], old["元のPDF"]):
        rows.append([w, "登録外国人統計", sub[v], int(n), f"古い表（{f}）"])
        if v == "特定活動":
            ref[w] = int(n)
    # 2006年末から：データ/全国_国籍_在留資格.csv（2006〜2011年末は内わけもここから）
    nat = pd.read_csv(os.path.join(DATA, "全国_国籍_在留資格.csv"))
    tk = nat[nat["在留資格"].str.startswith("特定活動")].groupby(["時点", "統計", "在留資格"])["人数"].sum()
    stat = {}
    for (w, s, v), n in tk.items():
        stat[w] = s
        if v in ("特定活動計", "特定活動"):
            ref[w] = int(n)
        if w < "2012":
            rows.append([w, s, TOTAL if v == "特定活動計" else v.replace("特定活動", ""), int(n), "データ/全国_国籍_在留資格.csv"])
    # 2012年12月末から：e-Stat の表
    for w, (kind, f) in sorted(sources().items()):
        path = os.path.join(SRC, f)
        if kind == "01-1":
            japan, by_nat = read_wide(path), {}
        else:
            japan, by_nat = read_long(path, kind)
        rows += [[w, stat.get(w, "在留外国人統計"), p, v, f] for p, v in japan.items()]
        if w >= "2022-12":
            nat_rows += [[w, n, p, v, f] for n, d in sorted(by_nat.items()) for p, v in d.items() if v and p != TOTAL]
            for p in japan:  # 国籍ごとの合計＝全国の数
                s = sum(d.get(p, 0) for d in by_nat.values())
                if s != japan[p]:
                    CHECKS.append([w, f"国籍の合計（{p}）", s, japan[p], s - japan[p]])
    # 確かめ：在留目的の合計＝総数、総数＝ほかの表の数
    byw = defaultdict(dict)
    for w, _, p, v, _ in rows:
        byw[w][p] = v
    for w, d in sorted(byw.items()):
        s = sum(v for p, v in d.items() if p != TOTAL and not p.startswith(OUTSIDE))
        CHECKS.append([w, "在留目的の合計＝総数", s, d[TOTAL], s - d[TOTAL]])
        if w in ref:
            CHECKS.append([w, "総数＝在留資格の表の特定活動", d[TOTAL], ref[w], d[TOTAL] - ref[w]])
    # 新しく入国した人（1年間）
    entry_rows = []
    for r in csv.DictReader(open(ENTRY_LIST, encoding="utf-8-sig")):
        f = "出入国管理統計/" + r["ファイル"]
        got = read_entry(os.path.join(SRC, f))
        entry_rows += [[r["年"], p, no, v, f] for p, no, v in got]
        total = next(v for p, _, v in got if p == TOTAL)
        s = sum(v for p, _, v in got if p != TOTAL)
        CHECKS.append([r["年"], "新しく入国した人：入国目的の合計＝総数", s, total, s - total])
    pd.DataFrame(rows, columns=["時点", "統計", "在留目的", "人数", "出どころ"]).to_csv(os.path.join(OUT, "在留目的別.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(nat_rows, columns=["時点", "国籍・地域", "在留目的", "人数", "出どころ"]).to_csv(os.path.join(OUT, "国籍_在留目的.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(entry_rows, columns=["年", "入国目的", "告示の号", "人数", "出どころ"]).to_csv(os.path.join(OUT, "新規入国_入国目的.csv"), index=False, encoding="utf-8-sig")
    pd.DataFrame(CHECKS, columns=["時点", "確かめ", "数", "比べる数", "差"]).to_csv(os.path.join(OUT, "確かめ.csv"), index=False, encoding="utf-8-sig")
    bad = [c for c in CHECKS if c[4]]
    print(f"在留目的別 {len(rows)} 行（{len(byw)} 時点）、国籍_在留目的 {len(nat_rows)} 行、新規入国_入国目的 {len(entry_rows)} 行")
    print("確かめ：" + ("すべて合っています" if not bad else f"合わない所 {len(bad)}"))
    for c in bad:
        print("  ", c)


if __name__ == "__main__":
    main()
