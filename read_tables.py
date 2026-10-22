# e-Stat の在留外国人統計の Excel を、1行に1つの数が入った形（縦長の表）に読み出すための部品
#   ・横に並ぶ表（第1表 国籍×在留資格、第4表 都道府県×国籍、第5表 都道府県×在留資格 など）… read_wide
#   ・市区町村の表（第7表など。都道府県・市・区・郡・町村が3つの列に段になって並ぶ）     … read_municipal
#   ・テーブルデータ（1行ずつのデータがシートに入っている 2016〜2023年ごろの形）          … read_rows_sheet
#   名前のそろえ方（空白・番号を取る、「青森」→「青森県」など）もここ
import re
import zipfile
from collections import defaultdict

import pandas as pd

PREFS = ("北海道 青森県 岩手県 宮城県 秋田県 山形県 福島県 茨城県 栃木県 群馬県 埼玉県 千葉県 東京都 神奈川県 新潟県 富山県 石川県 福井県 "
         "山梨県 長野県 岐阜県 静岡県 愛知県 三重県 滋賀県 京都府 大阪府 兵庫県 奈良県 和歌山県 鳥取県 島根県 岡山県 広島県 山口県 徳島県 "
         "香川県 愛媛県 高知県 福岡県 佐賀県 長崎県 熊本県 大分県 宮崎県 鹿児島県 沖縄県").split()
UNKNOWN_PREF = "未定・不詳"
# 国籍・地域の表にある州の小計（国ではない）
REGIONS = {"アジア", "ヨーロッパ", "アフリカ", "北米", "南米", "北アメリカ", "南アメリカ", "オセアニア", "無国籍の計"}
TOTAL = "総数"
TOTAL_ROW = "合計の行"  # 市区町村の表の、政令市（区の合計）と東京の特別区（23区の合計）の行
GUN = re.compile(r"郡(\(.+\)|（.+）)?$")


def clean(v):
    """空白・改行・全角空白を取る。数字だけのセル（Excel の数）はそのまま文字に"""
    if v is None or (isinstance(v, float) and pd.isna(v)):
        return ""
    return re.sub(r"[\s　]", "", str(v))


def pref_name(v):
    """「13：東京都」「北   海   道」「青森」「不詳・不定」などを「青森県」「未定・不詳」の形に。都道府県でなければ None"""
    s = clean(v).split("：")[-1]
    if s in PREFS:
        return s
    if s.startswith(("未定", "不詳")) or s in ("不詳・不定", "未定・不祥"):  # 2012年の第5表は「未定・不祥」
        return UNKNOWN_PREF
    hit = [p for p in PREFS if p[:-1] == s]
    return hit[0] if hit else None


def nat_name(v):
    """「01_008：インドネシア」「02アジア008インドネシア」「インドネシア」を国・地域の名前だけに"""
    s = clean(v).split("：")[-1]
    return re.sub(r"^\d{2}\D*?\d{3}", "", s)


def visa_name(v):
    """「03：教授」「教　　授」などを在留資格の名前だけに（「技術・人文知識・\n国際業務」の改行も取る）"""
    return clean(v).split("：")[-1]


def to_int(v):
    s = clean(v).replace(",", "")
    if s in ("", "-", "－", "―", "‐", "…", "x", "X", "*"):
        return 0
    return int(float(s))


def _merged_spans(path, sheet, row):
    """Excel（.xlsx）の、ある行（0から数える）でつながったマスの列の範囲 [(はじめ, おわり), …]"""
    if not str(path).lower().endswith(".xlsx"):
        return []
    from openpyxl import load_workbook
    wb = load_workbook(path)
    ws = wb.worksheets[sheet] if isinstance(sheet, int) else wb[sheet]
    spans = [(r.min_col - 1, r.max_col - 1) for r in ws.merged_cells.ranges if r.min_row == row + 1 and r.max_col > r.min_col]
    # 「選択範囲内で中央」の見出し（2020年12月末の第1表の「特定活動」）：名前のあるマスから、空で同じ揃え方のマスが続く所まで
    cells = list(ws.iter_rows(min_row=row + 1, max_row=row + 1))[0]
    k = 0
    while k < len(cells):
        if cells[k].value is not None and cells[k].alignment.horizontal == "centerContinuous":
            e = k
            while e + 1 < len(cells) and cells[e + 1].value is None and cells[e + 1].alignment.horizontal == "centerContinuous":
                e += 1
            if e > k:
                spans.append((k, e))
            k = e + 1
        else:
            k += 1
    return spans


def _header(v, path=None, sheet=0):
    """「総数」のある行を見出しとみなし、その列より左を行の名前の列にする。
    見出しが2段のとき（上の段に「技能実習」、下の段に「１号イ」など）は上の段の名前を前に付ける"""
    for i in range(min(25, len(v))):
        row = [clean(x) for x in v.iloc[i]]
        if TOTAL in row:
            tc = row.index(TOTAL)
            names = row[:]
            # すぐ上の段にだけ名前のある列（2020年12月末の第1表：「高度専門職1号イ」「技能実習1号イ」などは総数より1段上）。
            # 上の段の名前の下の「計」は「特定活動計」「日本人の配偶者等計」に（「計」が2つあると見分けられない）
            if i > 0:
                above = [clean(x) for x in v.iloc[i - 1]]
                parents = False
                for j in range(tc + 1, len(names)):
                    if j < len(above) and above[j]:
                        if not names[j]:
                            names[j] = above[j]
                        elif names[j] == "計":
                            names[j] = above[j] + "計"
                            parents = True
                # 上の段の名前が横につながったマス（「特定活動」）の下の細かい区分は、その名前を前に付ける（「その他」→「特定活動その他」）
                if parents and path:
                    for a, b in _merged_spans(path, sheet, i - 1):
                        if a < len(above) and above[a]:
                            for k in range(a + 1, min(b, len(names) - 1) + 1):
                                if names[k] and not names[k].startswith(above[a]):
                                    names[k] = above[a] + names[k]
            # すぐ下の段にも名前があれば（総数の列は空）、2段の見出し
            if i + 1 < len(v):
                below = [clean(x) for x in v.iloc[i + 1]]
                if below[tc] == "" and any(below[j] and not re.fullmatch(r"[\d,.]+", below[j]) for j in range(tc + 1, len(below))):
                    top = ""
                    for j in range(tc + 1, len(names)):
                        top = names[j] or top
                        if below[j]:
                            names[j] = top + below[j] if top and not below[j].startswith(top) else below[j]
                    return i + 1, tc, names
            return i, tc, names
    raise ValueError("見出しの行（総数）が見つかりません")


def read_wide(path, sheet=0):
    """横に並ぶ表を (行の名前, 段の深さ, 列の名前, 人数) の DataFrame に。総数の列は 列='総数'。
    段の深さ: 行の名前が左から何番目の列にあったか（0＝いちばん左。州の小計・総数の行は 0 のことが多い）"""
    v = pd.read_excel(path, header=None, sheet_name=sheet)
    hr, tc, names = _header(v, path, sheet)
    out = []
    for i in range(hr + 1, len(v)):
        labels = [clean(x) for x in v.iloc[i, :tc]]
        lv = max((k for k, x in enumerate(labels) if x), default=None)
        if lv is None:
            continue
        label = labels[lv]
        if label.startswith(("注", "(注", "（注", "※", "資料")):
            break
        vals = v.iloc[i, tc:]
        if all(clean(x) == "" for x in vals):
            continue
        for j in range(tc, v.shape[1]):
            col = names[j]
            if not col:
                continue
            out.append((label, lv, col, to_int(v.iloc[i, j])))
    return pd.DataFrame(out, columns=["行", "段", "列", "人数"])


def read_municipal(path, sheet=0):
    """市区町村の表を (都道府県, 郡, 市区町村, 種別, 列の名前, 人数) の DataFrame に。
    種別: 市区町村（ふつうの市町村・東京の特別区）／区（政令市の区。名前は「札幌市中央区」の形）／
          合計の行（政令市の区の合計、東京の「特別区」＝23区の合計）
    合計の行と区の行の両方を足すと二重になるので、合計するときは 種別 が「合計の行」の行を除く。
    表の形は2つ:
      2019年6月末まで … 都道府県と最初の市が同じ行。市・郡が2列目、区・町村が3列目。都道府県の合計の行はない
      2019年12月末から … 都道府県の合計の行があり、市・町村（「〇〇郡〇〇町」）が1列目、区が2列目
    町村の名前は「郡」を取って 郡 の列に入れる（年で書き方が違うため）"""
    v = pd.read_excel(path, header=None, sheet_name=sheet)
    hr, tc, names = _header(v)
    if "コード" in names[0]:
        return _read_municipal_coded(v, hr, tc, names)
    cols = [j for j in range(tc) if any(clean(v.iloc[i, j]) for i in range(hr + 1, len(v)))]  # 空の列は使わない
    out = []
    pref = gun = parent = None  # parent: 直前の数のある市（区が続けば政令市・特別区）
    city_col = None             # 市が何列目にあるか（表の形で違う）

    def add(name, kind, vals, g=None):
        for k, x in enumerate(vals):
            if names[tc + k]:
                out.append([pref, g or "", name, kind, names[tc + k], to_int(x)])

    for i in range(hr + 1, len(v)):
        labels = [clean(v.iloc[i, j]) for j in cols]
        filled = []
        for k, x in enumerate(labels):  # 同じ名前が2列に並ぶ形（2019年12月末の第7表の2）は1つに
            if x and x not in [y for _, y in filled]:
                filled.append((k, x))
        if not filled:
            continue
        k0, first = filled[0]
        if first.startswith(("注", "(注", "（注", "※", "資料")):
            break
        vals = [v.iloc[i, j] for j in range(tc, v.shape[1])]
        has_val = any(clean(x) != "" for x in vals)
        if first in ("総数", "全国"):
            continue
        p = pref_name(first)
        if p:
            pref, gun, parent = p, None, None
            if p == UNKNOWN_PREF:
                add(UNKNOWN_PREF, "市区町村", vals)
                continue
            if len(filled) == 1:  # 都道府県の合計の行（2019年12月末からの形）
                continue
            k, name = filled[1]
        else:
            k, name = k0, first
        if city_col is None:
            city_col = k
        if k == city_col:  # 市・郡・（2019年12月末からは）町村
            if GUN.search(name) or not has_val:  # 郡の行（数が空か0。「上川郡(石狩)」の形もある）。町村は次の行から
                gun, parent = name, None
                if any(to_int(x) for x in vals):  # 郡の行に数が入っている年がある（2016年6月末の八頭郡14人など）。消さずに残す
                    add(name + "（町村の記載なし）", "その他", vals, name)
                continue
            m = re.fullmatch(r"(.+?郡)(.+[町村])", name)
            if m:  # 「目梨郡羅臼町」の形
                add(m.group(2), "市区町村", vals, m.group(1))
                parent = None
                continue
            gun, parent = None, name
            add(name, "市区町村", vals)
        else:  # 政令市の区・東京の特別区、または郡の町村
            if parent is None:
                add(name, "市区町村", vals, gun)
                continue
            for row in reversed(out):  # 区が出たので、直前の市の行は「合計の行」
                if row[2] != parent:
                    break
                row[3] = TOTAL_ROW
            if parent.startswith("特別区"):  # 東京の23区は、それぞれが市区町村
                add(name, "市区町村", vals)
            else:
                if name.startswith("市") and name.endswith("区"):  # 2019年12月末・2020年6月末の表は相模原市の南区が「市南区」
                    name = name[1:]
                add(name if name.startswith(parent) else parent + name, "区", vals)
    df = pd.DataFrame(out, columns=["都道府県", "郡", "市区町村", "種別", "列", "人数"])
    df.insert(1, "市区町村コード", "")
    return df


def _read_municipal_coded(v, hr, tc, names):
    """2021年12月末からの形（1列目が市区町村コード、2列目が「札幌市中央区」のような名前）。
    コードの下3けたが 000 の行は都道府県の合計。次の行の名前がこの行の名前で始まれば政令市（合計の行）、東京は 13100（特別区部）"""
    rows = []
    for i in range(hr + 1, len(v)):
        code, name = clean(v.iloc[i, 0]), clean(v.iloc[i, 1])
        if not name or name in ("総数", "全国"):
            continue
        if name.startswith(("注", "(注", "（注", "※", "資料")):
            break
        rows.append((code.zfill(5) if code.isdigit() else code, name, [v.iloc[i, j] for j in range(tc, v.shape[1])]))
    out = []
    pref = parent = None
    for n, (code, name, vals) in enumerate(rows):
        p = pref_name(name)
        if p and (code.endswith("000") or p == UNKNOWN_PREF or not code):
            pref, parent = p, None
            if p != UNKNOWN_PREF:
                continue
            kind, gun = "市区町村", ""
        else:
            nxt = rows[n + 1][1] if n + 1 < len(rows) else ""
            gun = ""
            if code == "13100" or (nxt.startswith(name) and nxt != name and name.endswith("市")):
                kind, parent = TOTAL_ROW, (code, name)
            elif parent and code[:3] == parent[0][:3] and code != parent[0] and (name.startswith(parent[1]) or parent[0] == "13100"):
                kind = "市区町村" if parent[0] == "13100" else "区"
            else:
                parent = None
                kind = "市区町村"
                m = re.fullmatch(r"(.+?郡)(.+[町村])", name)
                if m:
                    gun, name = m.group(1), m.group(2)
        for k, x in enumerate(vals):
            if names[tc + k]:
                out.append([pref, code, gun, name, kind, names[tc + k], to_int(x)])
    return pd.DataFrame(out, columns=["都道府県", "市区町村コード", "郡", "市区町村", "種別", "列", "人数"])


def read_rows_sheet(path, keep):
    """テーブルデータ（いちばん大きいシートに1行ずつのデータ）を読み、keep の列ごとに人数を足した DataFrame に。
    年齢・性別の列は足し合わせて消える。openpyxl だと遅く重いので、シートの XML を直接読む"""
    z = zipfile.ZipFile(path)
    ss = [re.sub(r"<[^>]+>", "", re.sub(r"<rPh.*?</rPh>", "", m, flags=re.S))
          for m in re.findall(r"<si>(.*?)</si>", z.read("xl/sharedStrings.xml").decode("utf-8"), re.S)]
    sheet = max((i for i in z.infolist() if i.filename.startswith("xl/worksheets/sheet")), key=lambda i: i.file_size).filename
    data = z.read(sheet).decode("utf-8")
    cell = re.compile(r'<c r="([A-Z]+)\d+"([^>]*?)(?:/>|>(?:<f>[^<]*</f>)?(?:<v>([^<]*)</v>|<is><t>([^<]*)</t></is>)?</c>)')
    head, agg = None, defaultdict(int)
    for row in re.finditer(r"<row [^>]*>(.*?)</row>", data, re.S):
        vals = {}
        for col, attrs, v, inl in cell.findall(row.group(1)):
            vals[col] = ss[int(v)] if ' t="s"' in attrs and v else (inl or v or "")
        if head is None:
            head = {clean(v): k for k, v in vals.items()}
            num = next(k for k in head if "人数" in k)
            ck = [head[c] for c in keep]
            cn = head[num]
            continue
        agg[tuple(vals.get(c, "") for c in ck)] += to_int(vals.get(cn))
    return pd.DataFrame([(*k, n) for k, n in agg.items()], columns=[*keep, "人数"])


def read_long(path, sheet=0):
    """2020年12月末からの「1行に1つの数」の表（〇〇m.xlsx。列は 時点・都道府県・市区町村・国籍・地域 などと人数）。
    見出しの列の名前はそのまま（空白は取る）、人数の列は「人数」にそろえる。注の行は除く"""
    v = pd.read_excel(path, header=None, sheet_name=sheet)
    hr = next(i for i in range(15) if any("人数" in clean(x) for x in v.iloc[i]))
    head = [clean(x) for x in v.iloc[hr]]
    df = v.iloc[hr + 1:].copy()
    df.columns = head
    num = next(c for c in head if "人数" in c)
    df = df.rename(columns={num: "人数"})
    keys = [c for c in head if c not in ("時点", num) and c]
    for c in keys:
        df[c] = df[c].map(clean)
    df = df[~df[keys[0]].str.startswith(("注", "(注", "（注", "※", "資料"))]
    df = df[df["人数"].map(clean) != ""]
    df["人数"] = df["人数"].map(to_int)
    return df[keys + ["人数"]].reset_index(drop=True)


def read_powerpivot(path, keep):
    """2024年12月末からのテーブルデータ（Power Pivot）。keep の列ごとに人数を足した DataFrame に。
    pbixray は表の名前が R で始まると読み飛ばす（正規表現の書き違い）ので、ここで直してから使う（イスラム教系/統計/estat_pp.py と同じ直し方）"""
    from pbixray.meta import xml_source
    from pbixray.meta.xml_source import XmlMetadataSource

    def _fixed(self):
        self._load_xml_collection(
            re.compile(r"^(?!H\$|R\$)([^$]*?)\.(\d+)\.tbl\.xml$"),
            lambda s: xml_source.XMObjectDocument.from_xml_string(s).root_object,
            self._tbl_objects, "table",
        )

    XmlMetadataSource._extract_tbl_metadata = _fixed
    from pbixray import PBIXRay
    m = PBIXRay(path)
    t = list(m.tables)[0]
    cols = list(m.schema[m.schema["TableName"] == t]["ColumnName"])
    num = next(c for c in cols if "人数" in c)
    df = m.get_table(t, columns=[*keep, num])
    for c in keep:
        df[c] = df[c].astype(str).map(clean)
    df = df.rename(columns={num: "人数"})
    df["人数"] = pd.to_numeric(df["人数"], errors="coerce").fillna(0).astype(int)
    return df.groupby(keep, as_index=False, observed=True)["人数"].sum()


# ---------- 在留資格のまとめ方（サイトと、内わけの確かめで使う） ----------

# 細かい区分（国の表で年によって分けてあるもの）を、まとめた在留資格の名前に
SUB_VISA = {**{n: "特定活動" for n in ("アマスポーツ選手家族", "アマスポーツ選手本人", "インターンシップ", "ワーキング・ホリデー", "医療滞在・同伴者",
                                       "家事使用人", "本邦大卒者", "本邦大卒者の家族", "特定研究等及び情報処理家族", "特定研究等及び情報処理本人",
                                       "高度人材家族", "難民認定手続中", "ＥＰＡ対象者家族", "ＥＰＡ対象者本人", "人身取引等")},
            "短期観光": "短期滞在", "短期商用": "短期滞在", "滞在その他": "短期滞在", "滞在文化・学術活動": "短期滞在", "滞在親族訪問": "短期滞在",
            "日本人の子": "日本人の配偶者等", "日本人の配偶者": "日本人の配偶者等"}


def visa_group(name):
    """在留資格の見出しの名前と、小計の行かどうか。技能実習１号イ などの細かい区分は、まとめた名前に。
    小計の行（「〇〇合計」「〇〇計」）は (〇〇, True)。どれの小計か分からない「計」だけの行は (None, True)"""
    if name == "総数":
        return name, False
    sub = name.endswith("計")
    base = name[:-2] if name.endswith("合計") else (name[:-1] if sub else name)
    base = {"短期": "短期滞在"}.get(base, base)
    m = re.match(r"(技能実習|特定技能|高度専門職|特定活動|日本人の配偶者等)", base)
    base = m.group(1) if m else SUB_VISA.get(base, base)
    return (base or None), sub


def group_visas(df):
    """在留資格をまとめた名前で足す。小計の行がある在留資格は小計の行を使う（細かい区分が全部は載っていない年がある。2015年の全国の特定活動など）"""
    df = df.copy()
    g = df["在留資格"].map(visa_group)
    df["在留資格"], df["_sub"] = [b for b, _ in g], [s for _, s in g]
    df = df.dropna(subset=["在留資格"])
    keys = [c for c in df.columns if c not in ("人数", "_sub")]
    has_sub = df.groupby(keys, dropna=False)["_sub"].transform("max")
    df = df[df["_sub"] == has_sub]
    return df.groupby(keys, as_index=False, dropna=False)["人数"].sum()
